# from openpyxl import Workbook
# wb = Workbook() # создаем новую книгу
# ws = wb.active # выбираем активный лист
# ws.title = 'Студенты'
#
# students = [
#     ['№', 'Name student'],
#     [1, 'Ali'],
#     [2, 'Moldir'],
#     [3, 'Mihail']
# ]
#
# # добавляем данные в книгу
# for row in students:
#     ws.append(row)
#
# #сохраняем файл
# wb.save('students.xlsx')
# =====================================================================
from openpyxl import load_workbook

wb = load_workbook('students.xlsx')
ws = wb.active

# for row in ws.iter_rows(values_only=True):
#     print(row)

students_list = [list(row) for row in ws.iter_rows(values_only=True)]
print(students_list)
students_list[1][1] = 'Anastasia'
print(students_list)